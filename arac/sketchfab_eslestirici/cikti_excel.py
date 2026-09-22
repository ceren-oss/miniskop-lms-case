"""eslesmeler.xlsx uretimi (openpyxl)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .kazanimlar import Kazanim
from .sketchfab import Model
from .terimler import TerimKarari

BASLIKLAR = [
    ("Seviye", 12),
    ("Etkinlik", 26),
    ("Kazanım", 46),
    ("Sıra", 6),
    ("Model Adı", 34),
    ("Yazar", 20),
    ("Link", 38),
    ("Embed URL", 38),
    ("Beğeni", 9),
    ("Üçgen Sayısı", 13),
    ("Lisans", 24),
    ("Arama Terimi", 20),
    ("Puan", 8),
    ("Görüntülenme", 13),
    ("Animasyon", 10),
    ("Tablet Yükü", 12),
]

BASLIK_DOLGU = PatternFill("solid", fgColor="1F3B57")
BASLIK_YAZI = Font(color="FFFFFF", bold=True)
UYARI_DOLGU = PatternFill("solid", fgColor="FFF2CC")
LINK_YAZI = Font(color="0563C1", underline="single")


def _tablet_yuku(ucgen: int | None) -> str:
    if ucgen is None:
        return "bilinmiyor"
    if ucgen <= 150_000:
        return "hafif"
    if ucgen <= 500_000:
        return "orta"
    return "ağır"


def yaz(
    yol: Path,
    satirlar: list[tuple[Kazanim, list[Model]]],
    kararlar: dict[str, Any],
    ustbilgi: dict[str, Any],
    sahte_veri: bool = False,
) -> None:
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Eşleşmeler"

    satir_no = 1
    if sahte_veri:
        sayfa.cell(row=1, column=1, value=(
            "UYARI: Bu dosya --sahte-veri modunda üretildi. Satırlar GERÇEK Sketchfab "
            "modelleri değil, boru hattını denemek için üretilmiş örnek kayıtlardır."
        ))
        sayfa.cell(row=1, column=1).fill = UYARI_DOLGU
        sayfa.cell(row=1, column=1).font = Font(bold=True, color="9C5700")
        sayfa.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(BASLIKLAR))
        satir_no = 2

    baslik_satiri = satir_no
    for sutun, (baslik, genislik) in enumerate(BASLIKLAR, start=1):
        hucre = sayfa.cell(row=baslik_satiri, column=sutun, value=baslik)
        hucre.fill = BASLIK_DOLGU
        hucre.font = BASLIK_YAZI
        hucre.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        sayfa.column_dimensions[get_column_letter(sutun)].width = genislik

    satir_no = baslik_satiri + 1
    for kazanim, modeller in satirlar:
        for sira, model in enumerate(modeller, start=1):
            degerler = [
                kazanim.seviye,
                kazanim.etkinlik,
                kazanim.kazanim,
                sira,
                model.ad,
                model.yazar_ad,
                model.link,
                model.embed,
                model.begeni,
                model.ucgen,
                model.lisans,
                model.terim,
                model.puan,
                model.goruntulenme,
                model.animasyon,
                _tablet_yuku(model.ucgen),
            ]
            for sutun, deger in enumerate(degerler, start=1):
                hucre = sayfa.cell(row=satir_no, column=sutun, value=deger)
                if sutun == 3:
                    hucre.alignment = Alignment(wrap_text=True, vertical="top")
                elif sutun in (7, 8) and deger:
                    hucre.hyperlink = deger
                    hucre.font = LINK_YAZI
                elif sutun in (9, 10, 14, 15):
                    hucre.number_format = "#,##0"
                elif sutun == 13:
                    hucre.number_format = "0.000"
            satir_no += 1

    sayfa.freeze_panes = sayfa.cell(row=baslik_satiri + 1, column=1)
    if satir_no > baslik_satiri + 1:
        sayfa.auto_filter.ref = (
            f"A{baslik_satiri}:{get_column_letter(len(BASLIKLAR))}{satir_no - 1}"
        )

    _ozet_sayfasi(kitap, satirlar, kararlar, ustbilgi, sahte_veri)
    _eslesmeyen_sayfasi(kitap, satirlar, kararlar)
    _onerilmeyen_sayfasi(kitap, satirlar, kararlar)

    yol.parent.mkdir(parents=True, exist_ok=True)
    kitap.save(yol)


def _karar(kararlar: dict[str, Any], kazanim: Kazanim) -> TerimKarari:
    """TerimKarari'yi guvenle getirir (eksikse uygunsuz sayilir)."""
    return kararlar.get(kazanim.kimlik) or TerimKarari(uygun=False, neden="terim yok")


def _ozet_sayfasi(
    kitap: Workbook,
    satirlar: list[tuple[Kazanim, list[Model]]],
    kararlar: dict[str, Any],
    ustbilgi: dict[str, Any],
    sahte_veri: bool,
) -> None:
    sayfa = kitap.create_sheet("Özet")
    sayfa.column_dimensions["A"].width = 34
    sayfa.column_dimensions["B"].width = 52

    toplam_model = sum(len(m) for _, m in satirlar)
    eslesen = sum(1 for _, m in satirlar if m)
    onerilmeyen = sum(1 for k, _ in satirlar if not _karar(kararlar, k).uygun)
    hafif = sum(1 for _, ms in satirlar for m in ms if (m.ucgen or 0) <= 150_000)

    veriler: list[tuple[str, Any]] = [
        ("Çalıştırma", ""),
        ("Girdi dosyası", str(ustbilgi.get("girdi", ""))),
        ("Tarih", ustbilgi.get("tarih", "")),
        ("Terim üretimi", ustbilgi.get("terim_kaynagi", "")),
        ("Sketchfab isteği", ustbilgi.get("istek_sayisi", "")),
        ("Önbellek", ustbilgi.get("onbellek", "")),
        ("", ""),
        ("Sonuçlar", ""),
        ("Kazanım sayısı", len(satirlar)),
        ("3B model önerilen kazanım", eslesen),
        ("3B model uygun görülmeyen kazanım", onerilmeyen),
        ("Uygun ama model bulunamayan kazanım", len(satirlar) - eslesen - onerilmeyen),
        ("Toplam model satırı", toplam_model),
        ("Üretilen arama terimi",
         sum(len(_karar(kararlar, k).terimler) for k, _ in satirlar)),
        ("Tablet için hafif model (≤150k üçgen)", f"{hafif} / {toplam_model}"),
        ("", ""),
        ("Puanlama ağırlıkları", ""),
        ("Beğeni", ustbilgi.get("agirlik_begeni", "")),
        ("Üçgen sayısı", ustbilgi.get("agirlik_ucgen", "")),
        ("Gömülebilirlik", ustbilgi.get("agirlik_gomulebilirlik", "")),
    ]
    if sahte_veri:
        veriler.insert(0, ("!! ÖRNEK VERİ", "Gerçek Sketchfab sonucu değildir"))

    for satir, (etiket, deger) in enumerate(veriler, start=1):
        a = sayfa.cell(row=satir, column=1, value=etiket)
        sayfa.cell(row=satir, column=2, value=deger)
        if deger == "" and etiket:
            a.font = Font(bold=True)
        if etiket.startswith("!!"):
            a.font = Font(bold=True, color="9C5700")


def _eslesmeyen_sayfasi(
    kitap: Workbook,
    satirlar: list[tuple[Kazanim, list[Model]]],
    kararlar: dict[str, Any],
) -> None:
    """Uygun bulunup da Sketchfab'de karsiligi cikmayan kazanimlar."""
    eksikler = [
        (k, _karar(kararlar, k).terimler)
        for k, m in satirlar
        if not m and _karar(kararlar, k).uygun
    ]
    if not eksikler:
        return
    sayfa = kitap.create_sheet("Eşleşmeyenler")
    for sutun, (baslik, genislik) in enumerate(
        [("Seviye", 12), ("Etkinlik", 26), ("Kazanım", 60), ("Denenen Terimler", 40)], start=1
    ):
        hucre = sayfa.cell(row=1, column=sutun, value=baslik)
        hucre.fill = BASLIK_DOLGU
        hucre.font = BASLIK_YAZI
        sayfa.column_dimensions[get_column_letter(sutun)].width = genislik
    for satir, (kazanim, terim_listesi) in enumerate(eksikler, start=2):
        sayfa.cell(row=satir, column=1, value=kazanim.seviye)
        sayfa.cell(row=satir, column=2, value=kazanim.etkinlik)
        sayfa.cell(row=satir, column=3, value=kazanim.kazanim).alignment = Alignment(wrap_text=True)
        sayfa.cell(row=satir, column=4, value=", ".join(terim_listesi))


def _onerilmeyen_sayfasi(
    kitap: Workbook,
    satirlar: list[tuple[Kazanim, list[Model]]],
    kararlar: dict[str, Any],
) -> None:
    """3B modelin ogrenmeye katki saglamayacagi kazanimlar ve gerekceleri."""
    elenenler = [(k, _karar(kararlar, k)) for k, _ in satirlar
                 if not _karar(kararlar, k).uygun]
    if not elenenler:
        return
    sayfa = kitap.create_sheet("Model Önerilmeyenler")
    for sutun, (baslik, genislik) in enumerate(
        [("Seviye", 12), ("Etkinlik", 28), ("Kazanım", 58), ("Neden önerilmedi", 46)],
        start=1,
    ):
        hucre = sayfa.cell(row=1, column=sutun, value=baslik)
        hucre.fill = BASLIK_DOLGU
        hucre.font = BASLIK_YAZI
        sayfa.column_dimensions[get_column_letter(sutun)].width = genislik
    for satir, (kazanim, karar) in enumerate(elenenler, start=2):
        sayfa.cell(row=satir, column=1, value=kazanim.seviye)
        sayfa.cell(row=satir, column=2, value=kazanim.etkinlik)
        sayfa.cell(row=satir, column=3, value=kazanim.kazanim).alignment = Alignment(wrap_text=True)
        sayfa.cell(row=satir, column=4, value=karar.neden).alignment = Alignment(wrap_text=True)
    sayfa.freeze_panes = "A2"
