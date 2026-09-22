"""Arac icin regresyon testleri:  python3 -m pytest test_eslestirici.py -q"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from sketchfab_eslestirici import kazanimlar as kz
from sketchfab_eslestirici import terimler as trm
from sketchfab_eslestirici.cache import Onbellek
from sketchfab_eslestirici.cikti_html import _katla
from sketchfab_eslestirici.config import Ayarlar
from sketchfab_eslestirici.puanlama import (
    begeni_puani,
    elenmeli_mi,
    en_iyileri_sec,
    ucgen_puani,
)
from sketchfab_eslestirici.sahte import SahteIstemci, sahte_yanit
from sketchfab_eslestirici.sketchfab import SemaUyusmazligi, SketchfabIstemcisi


def ayarlar(tmp: Path, **fazla) -> Ayarlar:
    return Ayarlar(girdi=tmp / "x.csv", cikti_dizini=tmp, onbellek_dizini=tmp / "ob", **fazla)


# --------------------------------------------------------------- CSV okuma

def test_csv_noktali_virgul_ve_turkce_basliklar(tmp_path: Path):
    yol = tmp_path / "k.csv"
    yol.write_text("Seviye;Etkinlik Adı;Kazanım\n3. Sınıf;Gökyüzü;Ay'ı tanır.\n", encoding="utf-8")
    kayitlar, esleme = kz.oku(yol)
    assert len(kayitlar) == 1
    assert kayitlar[0].seviye == "3. Sınıf"
    assert kayitlar[0].kazanim == "Ay'ı tanır."
    assert esleme["kazanim"] == "Kazanım"


def test_csv_virgul_ve_ingilizce_basliklar(tmp_path: Path):
    yol = tmp_path / "k.csv"
    yol.write_text("grade,activity name,learning outcome\n5,Water,Explains the water cycle.\n",
                   encoding="utf-8")
    kayitlar, esleme = kz.oku(yol)
    assert kayitlar[0].etkinlik == "Water"
    assert esleme["seviye"] == "grade"


def test_csv_cp1254_ve_bos_satirlar(tmp_path: Path):
    yol = tmp_path / "k.csv"
    yol.write_bytes("Seviye;Etkinlik;Kazanım\n4;Işık;Işığı gözlemler.\n;;\n".encode("cp1254"))
    kayitlar, _ = kz.oku(yol)
    assert len(kayitlar) == 1          # bos satir atlanir
    assert kayitlar[0].kazanim == "Işığı gözlemler."


def test_csv_eksik_sutun_aciklayici_hata(tmp_path: Path):
    yol = tmp_path / "k.csv"
    yol.write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(ValueError, match="bulunamadi"):
        kz.oku(yol)


# ------------------------------------------------------------ terim yedegi

def test_sozluk_kelime_ortasinda_eslesmez():
    # "olusumlari" icindeki "ari" bal arisi olarak okunmamali
    terimler = trm.sozlukten_terimler("Volkanik oluşumları inceler.")
    assert "volcano" in terimler
    assert "honey bee" not in terimler


def test_sozluk_kisa_anahtar_tam_kelime_ister():
    assert "moon" not in trm.sozlukten_terimler("Ayrıca çevreyi korur.")
    assert "moon" in trm.sozlukten_terimler("Ay'ın evrelerini sıralar.")


def test_sozluk_en_fazla_dort_terim():
    uzun = "Güneş, Ay, Dünya, volkan, deprem, hücre, iskelet ve kalp konularını kapsar."
    assert len(trm.sozlukten_terimler(uzun)) <= 4


# --------------------------------------------------------------- puanlama

def test_begeni_puani_monoton_ve_sinirli():
    assert begeni_puani(0) == 0.0
    assert begeni_puani(10) < begeni_puani(100) < begeni_puani(1000)
    assert begeni_puani(10**9) <= 1.0


def test_ucgen_puani_agir_modeli_cezalandirir():
    assert ucgen_puani(20_000) == 1.0
    assert ucgen_puani(300_000) < ucgen_puani(100_000)
    assert ucgen_puani(5_000_000) == 0.0
    assert ucgen_puani(None) == 0.5      # bilgi yoksa notr


def test_puan_sifir_bir_araliginda(tmp_path: Path):
    a = ayarlar(tmp_path)
    modeller = SahteIstemci(Onbellek(tmp_path / "ob")).ara("volcano", 24)
    secilen, _ = en_iyileri_sec(modeller, a)
    assert secilen, "en az bir model secilmeli"
    assert all(0.0 <= m.puan <= 1.0 for m in secilen)
    assert [m.puan for m in secilen] == sorted((m.puan for m in secilen), reverse=True)


def test_yas_kisitli_ve_embedsiz_modeller_elenir(tmp_path: Path):
    a = ayarlar(tmp_path)
    modeller = SahteIstemci(Onbellek(tmp_path / "ob")).ara("volcano", 6)
    modeller[0].yas_kisitli = True
    modeller[1].embed = ""
    secilen, elenen = en_iyileri_sec(modeller, a)
    uidler = {m.uid for m in secilen}
    assert modeller[0].uid not in uidler
    assert modeller[1].uid not in uidler
    assert sum(elenen.values()) == 2


def test_ayni_model_iki_terimden_gelirse_tekillesir(tmp_path: Path):
    a = ayarlar(tmp_path, model_basina=5)
    istemci = SahteIstemci(Onbellek(tmp_path / "ob"))
    modeller = istemci.ara("volcano", 6) + istemci.ara("volcano", 6)
    secilen, _ = en_iyileri_sec(modeller, a)
    assert len({m.uid for m in secilen}) == len(secilen)


# --------------------------------------------------------------- onbellek

def test_onbellek_diskte_kalici(tmp_path: Path):
    ob = Onbellek(tmp_path / "ob")
    anahtar = Onbellek.anahtar("/search", {"q": "volcano"})
    assert ob.oku("sketchfab", anahtar) is None
    ob.yaz("sketchfab", anahtar, {"results": [1, 2]})
    yeni = Onbellek(tmp_path / "ob")          # sureci yeniden baslatmayi taklit et
    assert yeni.oku("sketchfab", anahtar) == {"results": [1, 2]}


def test_onbellek_bozuk_dosyayi_yutar(tmp_path: Path):
    ob = Onbellek(tmp_path / "ob")
    anahtar = Onbellek.anahtar("x")
    ob.yaz("sketchfab", anahtar, {"a": 1})
    (tmp_path / "ob" / "sketchfab" / f"{anahtar}.json").write_text("{bozuk", encoding="utf-8")
    assert ob.oku("sketchfab", anahtar) is None


def test_istek_onbellekten_dondurulur_ve_ag_kullanilmaz(tmp_path: Path, monkeypatch):
    ob = Onbellek(tmp_path / "ob")
    istemci = SketchfabIstemcisi(ob, bekleme=0.0)
    parametreler = {"type": "models", "q": "volcano", "sort_by": "-likeCount", "count": 3}
    ob.yaz("sketchfab", Onbellek.anahtar("/search", parametreler), {"results": []})

    def patlayan(*a, **k):
        raise AssertionError("onbellek varken ag istegi yapilmamali")

    monkeypatch.setattr(istemci.oturum, "get", patlayan)
    assert istemci._istek("/search", parametreler) == {"results": []}
    assert istemci.istek_sayisi == 0


# ------------------------------------------------------------------ sema

def test_sema_gercek_alan_adlarini_cozer(tmp_path: Path):
    istemci = SahteIstemci(Onbellek(tmp_path / "ob"))
    istemci.ara("volcano", 3)
    assert istemci.sema is not None and istemci.sema.saglikli
    assert istemci.sema.cozulen["begeni"] == "likeCount"
    assert istemci.sema.cozulen["ucgen"] == "faceCount"
    assert istemci.sema.cozulen["embed"] == "embedUrl"


def test_sema_results_yoksa_hata(tmp_path: Path):
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0)
    with pytest.raises(SemaUyusmazligi, match="results"):
        istemci._semayi_dogrula({"veriler": []})


def test_kati_sema_zorunlu_alan_eksikse_durdurur(tmp_path: Path):
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0, kati_sema=True)
    with pytest.raises(SemaUyusmazligi, match="zorunlu alanlar"):
        istemci._semayi_kur({"results": [{"identifier": "x", "title": "y"}]})


def test_sema_esnek_alan_adiyla_calisir(tmp_path: Path):
    """API 'likeCount' yerine 'likes' dondurse bile cozumleme surmeli."""
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0)
    govde = {"results": [{"uid": "a", "name": "n", "url": "u", "likes": 12}]}
    istemci._semayi_kur(govde)
    assert istemci.sema.cozulen["begeni"] == "likes"
    assert istemci.sema.cozulen["link"] == "url"


# -------------------------------------------------------------- html/misc

def test_html_katlama_python_ve_js_ile_uyumlu():
    assert _katla("Işık İSKELET") == "işik iskelet"
    assert _katla("IŞIK") == "işik"


def test_sahte_yanit_gercek_alan_adlarini_kullanir():
    govde = sahte_yanit("volcano", 2)
    ornek = govde["results"][0]
    for alan in ("uid", "name", "viewerUrl", "embedUrl", "likeCount", "faceCount",
                 "thumbnails", "user", "license"):
        assert alan in ornek


# ------------------------------------------------------- hiz siniri / geri cekilme

class SahteYanit:
    """requests.Response taklidi."""

    def __init__(self, kod: int, govde: dict | None = None, basliklar: dict | None = None):
        self.status_code = kod
        self._govde = govde or {}
        self.headers = basliklar or {}
        self.text = str(govde)

    def json(self):
        return self._govde

    def raise_for_status(self):
        if self.status_code >= 400:
            raise AssertionError(f"beklenmeyen durum {self.status_code}")


def test_429_retry_after_basligina_uyar(tmp_path: Path, monkeypatch):
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0)
    yanitlar = [
        SahteYanit(429, basliklar={"Retry-After": "3"}),
        SahteYanit(200, {"results": []}),
    ]
    beklemeler: list[float] = []
    monkeypatch.setattr("sketchfab_eslestirici.sketchfab.time.sleep", beklemeler.append)
    monkeypatch.setattr(istemci.oturum, "get", lambda *a, **k: yanitlar.pop(0))

    assert istemci._istek("/search", {"q": "x"}) == {"results": []}
    assert 3.0 in beklemeler          # sunucunun soyledigi sure beklendi
    assert istemci.istek_sayisi == 2


def test_sunucu_hatasi_ustel_geri_cekilir(tmp_path: Path, monkeypatch):
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0, azami_deneme=4)
    yanitlar = [SahteYanit(503), SahteYanit(503), SahteYanit(200, {"results": []})]
    beklemeler: list[float] = []
    monkeypatch.setattr("sketchfab_eslestirici.sketchfab.time.sleep", beklemeler.append)
    monkeypatch.setattr(istemci.oturum, "get", lambda *a, **k: yanitlar.pop(0))

    istemci._istek("/search", {"q": "x"})
    assert beklemeler == [2, 4]       # ustel artis


def test_tum_denemeler_tukenirse_hata(tmp_path: Path, monkeypatch):
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0, azami_deneme=2)
    monkeypatch.setattr("sketchfab_eslestirici.sketchfab.time.sleep", lambda s: None)
    monkeypatch.setattr(istemci.oturum, "get", lambda *a, **k: SahteYanit(500))
    with pytest.raises(RuntimeError, match="2 deneme basarisiz"):
        istemci._istek("/search", {"q": "x"})


def test_desteklenmeyen_filtre_dusurulup_tekrar_denenir(tmp_path: Path, monkeypatch):
    """API max_face_count'u reddederse arama filtresiz tekrarlanmali."""
    istemci = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0)
    gorulen: list[dict] = []

    def sahte_get(url, params=None, timeout=None):
        gorulen.append(dict(params or {}))
        if "max_face_count" in (params or {}):
            return SahteYanit(400, {"detail": "unknown parameter"})
        return SahteYanit(200, {"results": []})

    monkeypatch.setattr(istemci.oturum, "get", sahte_get)
    istemci.ara("volcano", 5, max_ucgen=100_000)
    assert "max_face_count" in gorulen[0]
    assert "max_face_count" not in gorulen[1]


def test_basarili_yanit_onbellege_yazilir(tmp_path: Path, monkeypatch):
    ob = Onbellek(tmp_path / "ob")
    istemci = SketchfabIstemcisi(ob, bekleme=0.0)
    monkeypatch.setattr(istemci.oturum, "get",
                        lambda *a, **k: SahteYanit(200, {"results": [{"uid": "a"}]}))
    parametreler = {"q": "x"}
    istemci._istek("/search", parametreler)
    # yeni surec: ag olmadan ayni sonuc gelmeli
    yeni = SketchfabIstemcisi(Onbellek(tmp_path / "ob"), bekleme=0.0)
    monkeypatch.setattr(yeni.oturum, "get", lambda *a, **k: pytest.fail("ag kullanilmamali"))
    assert yeni._istek("/search", parametreler) == {"results": [{"uid": "a"}]}


def test_sozluk_eslesme_yoksa_bos_doner():
    """Sözlükte karşılığı yoksa Türkçe kelime uydurmak yerine boş dönmeli."""
    terimler = trm.sozlukten_terimler(
        "FAB.6. Merak ettiği konular hakkında deneyler yapabilme."
    )
    assert terimler == []


def test_baglam_imzasi_etkinlik_adini_kapsar(tmp_path: Path):
    """Kazanım genel bir müfredat koduysa asıl sinyal etkinlik adıdır."""
    yol = tmp_path / "k.csv"
    yol.write_text(
        "Seviye;Ünite;Etkinlik;Kazanım;Etkinlik İçeriği\n"
        "4 Yaş;Gözlem;Köpüren Dinozor;FAB.1. Bilimsel gözlem yapabilme;Sirke ve karbonat\n",
        encoding="utf-8",
    )
    kayitlar, esleme = kz.oku(yol)
    assert esleme["unite"] == "Ünite" and esleme["icerik"] == "Etkinlik İçeriği"
    assert esleme["etkinlik"] == "Etkinlik"        # "Etkinlik İçeriği" ile karismamali
    assert "Köpüren Dinozor" in kayitlar[0].baglam_imzasi
    assert "dinosaur" in trm.sozlukten_terimler(kayitlar[0].baglam_imzasi)


def test_terimsiz_kazanim_eslesmeyen_olarak_raporlanir(tmp_path: Path):
    from sketchfab_eslestirici.cikti_excel import yaz as excel_yaz
    from openpyxl import load_workbook

    kayit = kz.Kazanim(sira=1, seviye="3 Yaş", etkinlik="X", kazanim="FAB.6. ...")
    yol = tmp_path / "e.xlsx"
    excel_yaz(yol, [(kayit, [])], {kayit.kimlik: []}, {}, False)
    kitap = load_workbook(yol)
    assert "Eşleşmeyenler" in kitap.sheetnames
    assert kitap["Eşleşmeyenler"].cell(row=2, column=2).value == "X"
